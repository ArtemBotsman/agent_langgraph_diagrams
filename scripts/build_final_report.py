"""Build the final NIR report as a visually verified Word document."""

# ruff: noqa: E501  # Long source-backed Russian paragraphs remain readable here.

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = (
    ROOT
    / "docs"
    / "final"
    / "deliverables"
    / "AgentLangGraph_technology_report_final_2026-09-11.docx"
)
ASSETS = ROOT / "docs" / "obsidian_vault" / "assets"

NAVY = "17365D"
BLUE = "2F75B5"
CYAN = "1AA6B7"
PALE_BLUE = "DDEBF7"
PALE_CYAN = "E2F0F2"
PALE_AMBER = "FFF2CC"
PALE_GREEN = "E2F0D9"
GRAY = "5B6573"
LIGHT = "F5F7FA"
WHITE = "FFFFFF"
RED = "C00000"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, *, bold: bool = False, color: str = "1F2937", size: int = 9) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_table(table, widths: list[float] | None = None) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    for cell in table.rows[0].cells:
        set_cell_shading(cell, NAVY)
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.bold = True


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    for idx, value in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], value, bold=True, color=WHITE)
    for row_idx, row in enumerate(rows, start=1):
        cells = table.add_row().cells
        fill = WHITE if row_idx % 2 else LIGHT
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)
            set_cell_shading(cells[idx], fill)
    style_table(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_field(paragraph, field_name: str) -> None:
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = field_name
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_1, instr_text, fld_char_2])


def add_page_number(section) -> None:
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("— ")
    run.font.name = "Arial"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(GRAY)
    add_field(p, "PAGE")
    run = p.add_run(" —")
    run.font.name = "Arial"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(GRAY)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(10 if level == 1 else 7)
    p.paragraph_format.space_after = Pt(5)


def add_body(doc: Document, text: str, *, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0.75)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.12
    if bold_prefix and text.startswith(bold_prefix):
        prefix = p.add_run(bold_prefix)
        prefix.bold = True
        p.add_run(text[len(bold_prefix) :])
    else:
        p.add_run(text)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.7)
        p.paragraph_format.first_line_indent = Cm(-0.35)
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_callout(doc: Document, title: str, text: str, fill: str = PALE_BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p = cell.add_paragraph(text)
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.runs[0].font.name = "Arial"
    p.runs[0].font.size = Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_figure(doc: Document, filename: str, caption: str, width: float = 6.5) -> None:
    path = ASSETS / filename
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(7)
    for run in cap.runs:
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string(GRAY)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("1F2937")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    for name, size, color in [
        ("Title", 25, NAVY),
        ("Subtitle", 13, GRAY),
        ("Heading 1", 16, NAVY),
        ("Heading 2", 12, BLUE),
        ("Heading 3", 10.5, CYAN),
    ]:
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name != "Subtitle"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")


def build() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.core_properties.author = "Боцман Артем"
    doc.core_properties.last_modified_by = "Боцман Артем"
    doc.core_properties.title = (
        "AgentLangGraph: генерация трассируемых activity-диаграмм"
    )
    configure_styles(doc)
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(1.7)
    add_page_number(section)

    # Cover
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(55)
    r = p.add_run("УНИВЕРСИТЕТ ИТМО")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor.from_string(NAVY)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(35)
    title.paragraph_format.space_after = Pt(12)
    r = title.add_run("Агентная генерация Use Case\nи activity-диаграмм по требованиям")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(25)
    r.font.color.rgb = RGBColor.from_string(NAVY)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(45)
    r = sub.add_run("с проверяемой трассировкой и ограниченным циклом исправления")
    r.font.name = "Arial"
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor.from_string(CYAN)

    add_callout(
        doc,
        "Научно-исследовательская работа",
        "Инженерный release candidate и методика экспериментальной оценки. "
        "Hidden-оценка и экспертная валидация вынесены в отдельный научный gate.",
        PALE_CYAN,
    )

    info = doc.add_table(rows=4, cols=2)
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, (label, value) in enumerate(
        [
            ("Исполнитель", "Боцман Артём"),
            ("Направление", "Агентные системы и генерация аналитических артефактов"),
            ("Основной фреймворк", "LangGraph"),
            ("Дата версии", "11 сентября 2026 года"),
        ]
    ):
        set_cell_text(info.rows[idx].cells[0], label, bold=True, color=NAVY)
        set_cell_text(info.rows[idx].cells[1], value)
        set_cell_shading(info.rows[idx].cells[0], PALE_BLUE)
        set_cell_shading(info.rows[idx].cells[1], WHITE)
    doc.add_paragraph()
    p = doc.add_paragraph("Санкт-Петербург, 2026")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(55)
    p.runs[0].font.color.rgb = RGBColor.from_string(GRAY)
    doc.add_page_break()

    add_heading(doc, "Аннотация", 1)
    add_body(
        doc,
        "Разработан воспроизводимый LangGraph pipeline, который принимает структурированный "
        "SpecificationReq, формирует Use Cases, user stories, system stories и activity-диаграммы, "
        "а затем проверяет структуру и двустороннюю трассировку. В системе работают два агента: "
        "Use Case Agent и Activity Diagram Agent. Корневой граф оркестрирует их, материализует "
        "TraceManifest и запускает evaluator. LLM-критика и ограниченное исправление дополняют "
        "детерминированные валидаторы, но не заменяют их.",
    )
    add_body(
        doc,
        "Технический benchmark candidate содержит 30 авторских синтетических кейсов: 20 development "
        "и 10 hidden. Валидация benchmark проходит для 30/30 кейсов, mutation-suite обнаруживает 15/15 "
        "заявленных классов дефектов. Rule-based baseline B0 прогнан на всех 20 DEV-кейсах по три раза. "
        "Получено 60/60 технических завершений, stability 1.0 и semantic composite 0.408 с case-level "
        "bootstrap 95% CI [0.356; 0.463]."
    )
    add_body(
        doc,
        "Первый повторный live pilot на DEV-001 сравнил B0, B1 one-shot и FULL. Затем на трёх "
        "одинаковых DEV-кейсах разной сложности сравнили пять конфигураций. После уточнения "
        "контракта критика FULL завершил 3/3 запуска, получил semantic composite 0.714 и trace F1 "
        "0.933. B1 one-shot и вариант без повторного исправления получили 0/3 сквозных завершений. "
        "Пилот показывает необходимость контролируемого исправления, но не доказывает преимущество "
        "критика: вариант без него также завершил 3/3 запуска."
    )
    add_callout(
        doc,
        "Статус результатов",
        "Инженерная часть воспроизводима локально. Полный платный DEV, два независимых экспертных "
        "review и однократный hidden-run ещё не выполнены и не подменяются синтетическими выводами.",
        PALE_AMBER,
    )

    add_heading(doc, "Ключевые термины", 2)
    add_table(
        doc,
        ["Термин", "Пояснение"],
        [
            ["Agent (агент)", "Компонент с ролью, входным состоянием, LLM-инструментами и правилами переходов."],
            ["Graph (граф)", "Исполняемая схема вершин и рёбер, задающая порядок обработки и ветвления."],
            ["Baseline (базовая линия)", "Упрощённый метод сравнения, относительно которого измеряется вклад решения."],
            ["Traceability (трассировка)", "Машинные связи от требования к UC, шагам сценария и элементам диаграммы."],
            ["Deterministic validator", "Не-LLM проверка однозначных структурных правил."],
            ["Bounded repair", "Исправление с заранее заданным максимальным числом повторов."],
        ],
        [1.7, 4.8],
    )
    doc.add_page_break()

    add_heading(doc, "1. Постановка задачи", 1)
    add_heading(doc, "1.1. Контекст и границы", 2)
    add_body(
        doc,
        "Текущая работа ограничена генерацией аналитико-архитектурных артефактов до написания кода. "
        "Исторический проект полного жизненного цикла разработки ПО служит контекстом, но не входит "
        "в текущий scope. Это устраняет смешение с агентами-кодерами, тестировщиками и deploy-агентами."
    )
    add_heading(doc, "1.2. Цель", 2)
    add_body(
        doc,
        "Цель работы состоит в разработке и исследовании агентного конвейера, который преобразует "
        "подготовленные функциональные и нефункциональные требования в структурированные Use Cases, "
        "stories и activity-диаграммы с машинно проверяемой трассировкой."
    )
    add_heading(doc, "1.3. Актуальность", 2)
    add_body(
        doc,
        "Статический анализ эффективно извлекает структуру из существующего кода, но не решает задачу "
        "интерпретации бизнес-требований до реализации. LLM способен выделять акторов, цели и сценарии, "
        "однако его свободный вывод не гарантирует ссылочную целостность. Поэтому исследуется гибрид: "
        "LLM выполняет смысловую работу, а типизированные контракты и детерминированные проверки создают "
        "контролируемый контур качества."
    )
    add_heading(doc, "1.4. Исследовательские вопросы", 2)
    add_bullets(
        doc,
        [
            "RQ1. Улучшает ли многостадийный pipeline качество по сравнению с one-shot при одинаковом входе, модели и выходном контракте?",
            "RQ2. Какой вклад дают semantic critic и bounded repair?",
            "RQ3. Обеспечивает ли typed trace model полноту и ссылочную целостность?",
            "RQ4. Как меняются качество, стабильность, задержка и расход токенов с ростом сложности требований?",
        ],
    )
    add_heading(doc, "1.5. Критерии завершения", 2)
    add_table(
        doc,
        ["Аспект", "Критерий"],
        [
            ["Pipeline", "Не менее 90% DEV завершаются E2E без ручного исправления."],
            ["Traceability", "100% activity elements имеют источник или явный unsupported status."],
            ["UC provenance", "Не менее 95% обязательных UC elements имеют FR source или классифицированное происхождение."],
            ["Validators", "100% подготовленных мутаций поддерживаемых классов обнаружены."],
            ["Reproducibility", "Повторный расчёт по сохранённым данным даёт идентичные метрики."],
        ],
        [1.7, 4.8],
    )

    doc.add_page_break()
    add_heading(doc, "2. Входные и выходные контракты", 1)
    add_heading(doc, "2.1. SpecificationReq", 2)
    add_table(
        doc,
        ["Поле", "Тип", "Назначение"],
        [
            ["project_task", "str", "Исходный запрос пользователя."],
            ["project_name", "str", "Название проектируемого приложения."],
            ["project_goal", "str", "Бизнес-цель."],
            ["project_description", "str", "Контекст и границы продукта."],
            ["functional_requirements", "list[str]", "Функциональные требования, далее атомизируемые в FR."],
            ["non_functional_requirements", "list[str]", "Нефункциональные ограничения и свойства качества."],
        ],
        [1.8, 1.0, 3.7],
    )
    add_heading(doc, "2.2. Результат одного запуска", 2)
    add_bullets(
        doc,
        [
            "типизированный набор Use Cases и акторов;",
            "человекочитаемые UC, user stories и system stories;",
            "типизированные activity-модели и Mermaid source;",
            "TraceManifest с прямыми и обратными ссылками;",
            "validation reports и evaluation report;",
            "run manifest, token/cost telemetry и SQLite checkpoints для FULL.",
        ],
    )
    add_callout(
        doc,
        "Почему Mermaid",
        "Mermaid выбран как воспроизводимый текстовый формат визуализации. Внутренняя typed activity model "
        "хранится отдельно, поэтому renderer можно заменить. Строгий профиль UML остаётся решением, "
        "которое должен утвердить руководитель до scientific freeze.",
    )

    doc.add_page_break()
    add_heading(doc, "3. Архитектура решения", 1)
    add_figure(doc, "01_two_agents_pipeline.png", "Рисунок 1. Полный pipeline и два исполняемых агента.", 6.6)
    add_heading(doc, "3.1. Агенты", 2)
    add_table(
        doc,
        ["Компонент", "Ответственность", "Основной выход"],
        [
            ["Use Case Agent", "Преобразует требования в акторов, UC, шаги сценария и stories.", "UseCaseSet, user/system stories"],
            ["Activity Diagram Agent", "Строит поведенческий граф по каждому принятому UC.", "ActivityModel, Mermaid"],
            ["Root orchestrator", "Управляет подграфами, объединяет результаты, строит trace и запускает evaluator.", "Result bundle"],
        ],
        [1.6, 3.2, 1.7],
    )
    add_callout(
        doc,
        "Количество агентов",
        "В проекте два агента. Generator, critic и repair являются внутренними ролями и вершинами "
        "подграфов, а не тремя дополнительными автономными агентами.",
        PALE_GREEN,
    )
    add_heading(doc, "3.2. Вершины и рёбра", 2)
    add_table(
        doc,
        ["Граф", "Вершин", "Рёбер", "Назначение"],
        [
            ["Root graph", "6", "7", "Оркестрация UC, activities, trace, evaluation и export."],
            ["UC subgraph", "9", "13", "Generation, schema, validation, critic, decision, repair/failure."],
            ["Activity subgraph", "10", "14", "Generation, validation, critic, bounded repair, Mermaid."],
            ["Итого", "25", "34", "Рабочая статическая структура release candidate."],
        ],
        [1.8, 0.7, 0.7, 3.3],
    )
    add_figure(doc, "02_internal_agent_loop.png", "Рисунок 2. Внутренний ограниченный цикл проверки и исправления.", 6.2)
    add_heading(doc, "3.3. Логика ребёр", 2)
    add_body(
        doc,
        "Обычные рёбра передают типизированное состояние следующей вершине. Условные рёбра читают "
        "validation status, critic verdict и число попыток. Результат направляется в accept, repair "
        "или controlled failure. Лимит max_repair_attempts исключает бесконечный цикл."
    )

    add_heading(doc, "4. Трассировка и детерминированная валидация", 1)
    add_body(
        doc,
        "Основная цепочка происхождения имеет вид: FR → requirement atom → UC → scenario step → "
        "activity node/edge. LLM выбирает типизированные source IDs, а Python детерминированно "
        "создаёт единый TraceManifest. Прямые ссылки отвечают на вопрос «что было порождено этим "
        "требованием», обратные — «из какого источника появился этот элемент»."
    )
    add_heading(doc, "4.1. Правила", 2)
    add_table(
        doc,
        ["Правило", "Проверка", "Зачем"],
        [
            ["Уникальные ID", "Нет повторяющихся FR, UC, step, node, edge ID.", "Предотвращает неоднозначные ссылки."],
            ["Типы концов связи", "Каждая связь соединяет допустимые типы сущностей.", "Исключает логически невозможные trace links."],
            ["Ссылочная целостность", "Source и target реально существуют.", "Выявляет dangling references."],
            ["Достижимость", "Все activity nodes достижимы от start.", "Исключает изолированные фрагменты."],
            ["Branch guards", "У ветвлений есть условия перехода.", "Делает сценарий интерпретируемым."],
            ["Provenance", "Unsupported/assumption помечены и объяснены.", "Не скрывает неподтверждённые элементы."],
        ],
        [1.5, 3.0, 2.0],
    )
    add_heading(doc, "4.2. Почему валидатор не заменяется критиком", 2)
    add_body(
        doc,
        "LLM-критик полезен для смысловой полноты и согласованности, но вероятностен. Он не даёт "
        "строгой гарантии уникальности ID, достижения 100% coverage или существования каждой ссылки. "
        "Детерминированный валидатор повторяет одну и ту же проверку с одинаковым результатом, поэтому "
        "является обязательным gate до принятия артефакта."
    )
    add_figure(doc, "10_validator_mutations.png", "Рисунок 3. Mutation-suite: 15 из 15 заявленных дефектов обнаружены.", 6.2)

    add_heading(doc, "5. Benchmark и методика оценки", 1)
    add_figure(doc, "04_benchmark_map.png", "Рисунок 4. Состав benchmark candidate v1.0.", 6.5)
    add_table(
        doc,
        ["Характеристика", "Значение"],
        [
            ["Всего кейсов", "30"],
            ["Разбиение", "20 DEV / 10 hidden"],
            ["Сложность", "8 simple / 13 medium / 9 hard"],
            ["Язык", "13 RU / 17 EN"],
            ["Gold", "Акторы, UC, milestones, branches, FR-to-UC links"],
            ["Эквивалентность", "Canonical slots, aliases, optional slots, graph-equivalence rules"],
        ],
        [2.0, 4.5],
    )
    add_heading(doc, "5.1. Метрики", 2)
    add_table(
        doc,
        ["Группа", "Метрика", "Интерпретация"],
        [
            ["Формат", "Schema / structural validity", "Прошёл ли результат типизированные контракты."],
            ["Trace", "Integrity, forward/reverse coverage", "Полнота и корректность происхождения."],
            ["Семантика", "Actor/UC/milestone/branch F1", "Совпадение смысловых gold slots с эквивалентностями."],
            ["Риск", "Hallucination proxy", "Доля неподтверждённого содержания."],
            ["Система", "E2E success", "Прошёл ли весь обязательный контур."],
            ["Цена", "Latency, calls, tokens, cost, repairs", "Ресурсная стоимость качества."],
            ["Устойчивость", "Multi-run stability", "Сходство результатов повторных запусков."],
        ],
        [1.3, 2.3, 2.9],
    )
    add_callout(
        doc,
        "Ограничение composite",
        "Semantic composite используется как dashboard metric. Он не заменяет отдельные проверки, "
        "E2E gate и экспертную оценку. Пилот показал, что B1 может иметь более высокий composite, "
        "но формально не пройти трассировку.",
        PALE_AMBER,
    )
    add_heading(doc, "5.2. Экспертная шкала", 2)
    add_table(
        doc,
        ["Балл", "Смысл"],
        [
            ["1", "Критически неверно: результат нельзя использовать без полной переработки."],
            ["2", "Существенные ошибки: пропущены ключевые акторы, цели или ветви."],
            ["3", "Частично корректно: основа верна, но нужны заметные исправления."],
            ["4", "Почти корректно: только локальные правки без изменения структуры."],
            ["5", "Корректно и пригодно к использованию в заявленной постановке."],
        ],
        [0.7, 5.8],
    )

    add_heading(doc, "6. Сравниваемые методы", 1)
    add_table(
        doc,
        ["Метод", "Вход", "LLM", "Critic / repair", "Роль в исследовании"],
        [
            ["B0_RULE", "SpecificationReq", "Нет", "Нет", "Детерминированный нижний ориентир."],
            ["B1_ONESHOT", "SpecificationReq", "Да, один вызов", "Нет", "Прямой baseline той же модели."],
            ["FULL", "SpecificationReq", "Да", "Да", "Основной двухагентный pipeline."],
            ["FULL_NO_CRITIC", "SpecificationReq", "Да", "Только validator/repair", "Вариант без отдельного компонента semantic critic."],
            ["FULL_NO_REPAIR", "SpecificationReq", "Да", "Critic, 0 repair", "Вариант без отдельного компонента bounded repair."],
            ["Pyreverse", "Python-код", "Нет", "Нет", "Adjacent baseline для code-to-UML."],
            ["GitDiagram / DeepWiki", "Репозиторий", "Да", "Сервисный контур", "Смежные repository-to-diagram системы."],
        ],
        [1.3, 1.3, 1.0, 1.3, 1.7],
    )
    add_heading(doc, "7. Экспериментальные результаты", 1)
    add_heading(doc, "7.1. Repeated live pilot DEV-001", 2)
    add_table(
        doc,
        ["Метод", "E2E", "Semantic", "Activity trace", "Tokens/run", "Repairs/run", "n"],
        [
            ["B0", "2/2", "0.633", "1.000", "0", "0", "2"],
            ["B1", "0/2", "0.724", "0.881", "8 969", "0", "2"],
            ["FULL", "2/2", "0.670", "1.000", "33 445", "2", "2"],
        ],
        [1.0, 0.7, 0.9, 1.1, 1.0, 1.0, 0.4],
    )
    add_body(
        doc,
        "B1 не прошёл оба запуска из-за непротрассированных activity edges. Error analysis выявил "
        "8 случаев activity_edge_untraced и 2 случая activity_trace_coverage_below_threshold. FULL "
        "обнаружил четыре промежуточных activity_edge_untraced и исправил их; в финальном результате "
        "неразрешённых blocking issues не осталось."
    )
    add_figure(doc, "07_live_AD-UC001.png", "Рисунок 5. Пример сгенерированной activity-диаграммы DEV-001, UC-001.", 3.9)
    add_figure(doc, "08_live_AD-UC002.png", "Рисунок 6. Пример сгенерированной activity-диаграммы DEV-001, UC-002.", 4.8)
    add_heading(doc, "7.2. Полный B0 на DEV20", 2)
    add_table(
        doc,
        ["Группа", "Кейсов", "Запусков", "E2E", "Semantic", "95% CI", "Hallucination"],
        [
            ["Все DEV", "20", "60", "1.000", "0.408", "[0.356; 0.463]", "0.493"],
            ["Simple", "6", "18", "1.000", "0.543", "[0.481; 0.601]", "0.396"],
            ["Medium", "9", "27", "1.000", "0.359", "[0.293; 0.433]", "0.522"],
            ["Hard", "5", "15", "1.000", "0.335", "[0.296; 0.372]", "0.556"],
        ],
        [1.0, 0.7, 0.7, 0.7, 0.9, 1.1, 1.0],
    )
    add_figure(doc, "13_b0_dev20_quality.png", "Рисунок 7. Качество B0 по сложности требований.", 6.2)
    add_body(
        doc,
        "Все 60 запусков завершились технически, stability равна 1.0. Semantic composite ниже 0.5 "
        "получен у 15 из 20 уникальных кейсов; из-за трёх одинаковых повторов это соответствует "
        "45 из 60 запусков. Hallucination proxy выше 0.4 получен в 48 из 60 запусков. Следовательно, B0 подтверждает "
        "работоспособность инфраструктуры и воспроизводимость, но не решает смысловую задачу."
    )

    add_heading(doc, "7.3. Сравнение пяти конфигураций на трёх DEV-кейсах", 2)
    add_body(
        doc,
        "В пилот вошли B1-DEV-002, B1-DEV-010 и B1-DEV-020: простой русский, средний русский и "
        "сложный английский кейсы. Для каждой конфигурации выполнено по одному запуску на кейс. "
        "Все LLM-конфигурации использовали deepseek-flash, temperature=0, один endpoint и "
        "одинаковые входы; hidden не открывался."
    )
    add_table(
        doc,
        ["Метод", "E2E", "Semantic", "Branch F1", "Trace F1", "Tokens/run", "USD/run"],
        [
            ["B0", "3/3", "0.524", "0.000", "0.857", "0", "0.0000"],
            ["B1", "0/3", "0.487", "0.556", "0.619", "10 358", "0.0046"],
            ["FULL", "3/3", "0.714", "0.690", "0.933", "52 947", "0.0294"],
            ["Без критика", "3/3", "0.699", "0.500", "0.952", "37 011", "0.0238"],
            ["Без исправления", "0/3", "0.730", "0.786", "0.952", "9 380", "0.0065"],
        ],
        [1.25, 0.6, 0.9, 0.9, 0.85, 1.0, 0.8],
    )
    add_figure(
        doc,
        "20_method_comparison_dev3.png",
        "Рисунок 8. Прямое сравнение пяти конфигураций на одинаковых DEV-входах.",
        6.5,
    )
    add_body(
        doc,
        "B0 проходит структурный барьер, но не восстанавливает ветвления. Вариант без исправления "
        "получил высокий автоматический содержательный балл, но не сформировал принимаемый полный "
        "пакет. Поэтому semantic composite нельзя использовать без E2E gate. Прямой one-shot "
        "оказался примерно в 6.4 раза дешевле FULL, но не дал ни одного принимаемого результата на "
        "выбранных кейсах."
    )
    add_heading(doc, "7.4. Исправление контракта критика", 2)
    add_body(
        doc,
        "Первоначальный FULL завершил только 1/3 запусков: критик мог вводить новый критерий после "
        "каждого исправления. Теперь блокирующее замечание обязано ссылаться на конкретное требование "
        "и элемент модели, а замечания к стилю не блокируют результат. Завершение выросло до 3/3, "
        "но средняя стоимость - с 0.0239 до 0.0294 USD. Это подтверждает устранение инженерной "
        "ошибки, но не является статистическим доказательством полезности критика."
    )
    add_figure(
        doc,
        "21_critic_before_after.png",
        "Рисунок 9. Работа FULL до и после уточнения доказательного контракта критика.",
        6.4,
    )

    add_heading(doc, "8. Воспроизводимость, логи и безопасность", 1)
    add_heading(doc, "8.1. Сохраняемые доказательства", 2)
    add_bullets(
        doc,
        [
            "config.json с режимом и безопасной конфигурацией;",
            "generated specification, UC, stories, activity JSON и Mermaid;",
            "TraceManifest, validation/evaluation reports;",
            "calls telemetry, latency, token usage, repair count;",
            "SQLite checkpoints для возобновления FULL;",
            "results.json/results.csv, SHA-256 и offline reproducibility report.",
        ],
    )
    add_heading(doc, "8.2. Сетевые ошибки и budget guards", 2)
    add_body(
        doc,
        "OpenAI-compatible adapter повторяет только временные ошибки 408, 409, 425, 429 и 5xx, "
        "учитывает Retry-After и ограничивает задержку. Ошибки 400, 401 и 403 не повторяются. Перед "
        "live experiment задаются максимальные calls, tokens и estimated cost."
    )
    add_heading(doc, "8.3. Raw evidence", 2)
    add_body(
        doc,
        "Публичный result bundle хранит хэши prompt/response и sanitized telemetry. Точные request/response "
        "можно сохранить только явно через --capture-raw-private-dir. Приватный каталог получает права "
        "700, файлы — 600, API key не записывается, а путь не попадает в публичный config. Каталог "
        ".private_experiment_evidence исключён из Git."
    )
    add_heading(doc, "8.4. Повторный расчёт", 2)
    add_body(
        doc,
        "Offline verifier заново вычисляет summary, stability и SHA-256 по сохранённым bundles. Для "
        "B0 DEV20 подтверждены 60 ожидаемых и 60 полных bundles. Для repeated DEV-001 подтверждены "
        "6 из 6. Исторически сохранённые значения совпали; новые производные поля отчётливо отмечены "
        "как дополнительные, а не как расхождение."
    )

    doc.add_page_break()
    add_heading(doc, "9. Готовность по плану НИР", 1)
    add_table(
        doc,
        ["Этап", "Формальная готовность", "Готово", "Внешний или платный gate"],
        [
            ["Планирование", "95%", "План, RQ, риски, контрольные точки.", "Подтвердить сроки и план."],
            ["Анализ и проектирование", "95%", "Контракты, 2 агента, 25 вершин, 34 ребра.", "Утвердить UML/trace policy."],
            ["Benchmark и методика", "90%", "30 кейсов, gold, split, freeze, evaluator.", "2 эксперта и разрешение публикации."],
            ["Pipeline", "85%", "FULL, B0/B1, component variants, SQLite, retry, CLI.", "Полный paid FULL DEV20 и ≥90%."],
            ["Трассировка", "90%", "TraceManifest и mutation 15/15.", "Согласование правил и full DEV."],
            ["Интеграция", "80%", "RC, B0 DEV20×3, пять вариантов на DEV-3, verifier.", "Повторный FULL DEV20."],
            ["Эксперименты", "55%", "B0 full и пять вариантов на трёх DEV-кейсах.", "Повторы DEV20, 2 эксперта, hidden."],
            ["Финализация", "95%", "README, CI, report, deck, speech, demo.", "Репетиция и публикационное решение."],
        ],
        [1.3, 1.0, 2.4, 2.2],
    )
    add_callout(
        doc,
        "Что означает «доделать полностью»",
        "Код, воспроизводимость и материалы защиты можно завершить автономно. Научный DoD нельзя "
        "объявить 100%, пока реальные эксперты не поставили оценки, руководитель не утвердил правила, "
        "а платный repeated DEV и запечатанный hidden-run не были выполнены по согласованному протоколу.",
        PALE_AMBER,
    )

    add_heading(doc, "10. Ограничения валидности", 1)
    add_bullets(
        doc,
        [
            "Benchmark создан автором и имеет technical author freeze, но ещё не прошёл два независимых review.",
            "Прямое сравнение пяти конфигураций охватывает три DEV-кейса и один запуск на кейс.",
            "Semantic matcher и его порог не откалиброваны по экспертам.",
            "Alias модели провайдера не гарантирует immutable model revision.",
            "Порог 90% требует полного FULL DEV20, обобщаемость — однократного hidden-run после freeze.",
            "НФТ сохраняются и проверяются на уровне контракта/происхождения, но не все НФТ визуализируются activity-диаграммой.",
        ],
    )

    add_heading(doc, "11. План завершения эксперимента", 1)
    add_table(
        doc,
        ["Шаг", "Действие", "Условие перехода"],
        [
            ["1", "Утвердить MODEL_ID/API_BASE, UML profile, trace policy, benchmark publication.", "Письменное решение руководителя."],
            ["2", "Повторить FULL и варианты без отдельных компонентов на DEV20.", "Нет блокирующих pipeline/evaluator ошибок."],
            ["3", "Рассчитать доверительные интервалы, устойчивость и анализ классов ошибок.", "Метрики пересчитываются из сохранённых bundles."],
            ["4", "Два независимых expert review, agreement и scientific freeze.", "Зафиксированы commit/tag/hash."],
            ["5", "Однократный hidden-run без настройки по результату.", "Scientific freeze завершён."],
            ["6", "Финальная интерпретация, репетиция и публикация разрешённой части.", "Все числа восстанавливаются из bundles."],
        ],
        [0.5, 4.0, 2.0],
    )

    add_heading(doc, "12. Вывод", 1)
    add_body(
        doc,
        "Разработан полный исполняемый контур от SpecificationReq до трассируемых activity-диаграмм "
        "и отчёта качества. Детерминированные проверки обнаруживают структурные дефекты, а ограниченный "
        "agentic repair способен исправить непротрассированные элементы в live pilot. B0 доказывает "
        "воспроизводимость методики, но показывает недостаточность простых правил для смысловой задачи."
    )
    add_body(
        doc,
        "Инженерный release candidate готов к демонстрации. Пилот показывает вклад ограниченного "
        "исправления и цену многостадийного контура, но не позволяет объявить статистическое "
        "превосходство FULL или полезность критика. Научное завершение требует согласования моделей "
        "и правил, двух независимых экспертов, повторного FULL DEV20 и единственного hidden-run после "
        "научной фиксации."
    )

    doc.add_page_break()
    add_heading(doc, "Литература и источники", 1)
    sources = [
        "Hugging Face Agents Course. Unit 0 и LangGraph unit. https://huggingface.co/learn/agents-course/unit0/introduction",
        "LangGraph documentation. Overview, persistence and graph design. https://docs.langchain.com/oss/python/langgraph/overview",
        "ISO/IEC/IEEE 29148:2024. Requirements engineering. https://www.iso.org/standard/72089.html",
        "OMG UML 2.5.1 specification. https://www.omg.org/spec/UML/2.5.1/PDF",
        "Automated Generation of SysML Activity Diagrams from Industrial Requirements Using LLMs. RE 2026. https://conf.researchr.org/details/RE-2026/RE-2026-industrial-innovation-papers/4/",
        "R2ABench: requirements-to-architecture evaluation. 2026. https://arxiv.org/abs/2604.06683",
        "TraceLLM benchmark. 2026. https://arxiv.org/abs/2602.01253",
        "Systematic mapping of LLMs in model-driven engineering. 2026. https://link.springer.com/article/10.1007/s10664-026-10921-4",
        "Text2UML evaluation artifacts and scenarios. 2025. https://github.com/IlKaiser/text2uml",
        "Position bias in LLM-as-a-judge. 2025. https://aclanthology.org/2025.ijcnlp-long.18/",
    ]
    for idx, source in enumerate(sources, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        p.paragraph_format.space_after = Pt(5)
        p.add_run(f"{idx}. {source}")

    add_heading(doc, "Приложение A. Проверяемые артефакты", 1)
    add_table(
        doc,
        ["Артефакт", "Расположение в репозитории"],
        [
            ["Benchmark", "benchmark/v1_0_synthetic/cases.json"],
            ["Freeze record", "benchmark/v1_0_synthetic/freeze_record.json"],
            ["Методика", "docs/benchmark/BENCHMARK_AND_METRICS_V0_1.md"],
            ["Экспертная рубрика", "benchmark/expert_review/RUBRIC.md"],
            ["System Design", "docs/architecture/SYSTEM_DESIGN_V0_1.md"],
            ["Traceability Model", "docs/architecture/TRACEABILITY_MODEL.md"],
            ["B0 DEV20", "artifacts/benchmark_runs/b0-dev20-r3-2026-09-09"],
            ["Repeated live pilot", "artifacts/benchmark_runs/repeated-b0-b1-full-dev001-2026-09-09"],
            ["Five-method DEV pilot", "artifacts/technology_project_evidence_2026-09-11"],
            ["Full critic v2 run", "artifacts/benchmark_runs/full-critic-v2-dev3-deepseek-flash-2026-09-11"],
            ["Comparison report", "docs/research/TECHNOLOGY_PROJECT_COMPARISON_2026_09_11.md"],
            ["Reproducibility verifier", "scripts/verify_saved_experiment.py"],
            ["Error analysis", "scripts/analyze_saved_experiment.py"],
        ],
        [2.2, 4.3],
    )

    # Keep pagination deterministic across Word/LibreOffice.
    for section in doc.sections:
        section.header_distance = Cm(0.7)
        section.footer_distance = Cm(0.7)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
